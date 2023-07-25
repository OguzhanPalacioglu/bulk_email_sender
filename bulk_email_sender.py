import smtplib
from email.mime.text import MIMEText
import openpyxl
import time

# SMTP sunucusu ayarları
# SMTP server settings
smtp_server = 'smtp.gmail.com'
smtp_port = 587
sender_email = 'example@gmail.com'
sender_password = 'your_password'

# İsim, soyisim, yaş ve şehir için yer tutucular içeren e-posta şablonu
# Email template with placeholders for name, surname, age, and city
email_template = """Hello {{name}} {{surname}},
Your Age : {{age}} 
Your City : {{city}}
Regards,
O.P"""

# E-posta gönderen fonksiyon
# Function to send an email
def send_email(to_address, subject, body):
    msg = MIMEText(body)
    msg['Subject'] = subject
    msg['From'] = sender_email
    msg['To'] = to_address

    try:
        with smtplib.SMTP(smtp_server, smtp_port) as server:
            server.starttls()
            server.login(sender_email, sender_password)
            server.sendmail(sender_email, to_address, msg.as_string())
        print(f"E-posta başarıyla gönderildi: {to_address}")
    except Exception as e:
        print(f"E-posta gönderilirken hata oluştu: {e}")

if __name__ == "__main__":
    # Excel dosyasını yükleyelim
    # Load the Excel file.
    workbook = openpyxl.load_workbook("recipients.xlsx")
    worksheet = workbook.active

    # Sütun başlıklarını alalım
    # Get the column headers.
    columns = [col.value.lower() for col in worksheet[1]]

    # Alıcı bilgilerini bir liste olarak toplayalım.
    # Collect recipient information as a list.
    recipients = []
    for row in worksheet.iter_rows(min_row=2):
        recipient_info = {columns[i]: cell.value for i, cell in enumerate(row)}
        recipients.append(recipient_info)

    # E-postaları gönderelim
    # Send the emails.
    for recipient in recipients:
        personalized_template = email_template.replace('{{name}}', recipient['name']) \
                                              .replace('{{surname}}', recipient['surname']) \
                                              .replace('{{age}}', str(recipient['age'])) \
                                              .replace('{{city}}', recipient['city'])

        subject = "Otomatik E-posta Denemesi"
        body = personalized_template

        send_email(recipient['email'], subject, body)

        # Gönderimler arasında 3 saniye bekleme süresi ekleyelim ki spam'e düşmesin
        # Let's add a delay of 3 seconds between each email to avoid being marked as spam.
        time.sleep(3)

    print("Program Sonlandı.")
